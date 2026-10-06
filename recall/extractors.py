"""Turn every supported file type into normalized Markdown.

Embedded images (PDF/DOCX/PPTX) are written to a cache directory and referenced as
``/api/cache-image/<name>``; images referenced from Markdown/HTML notes are rewritten to
``/api/file?path=<path relative to the notes root>``. Links between notes become ``#doc=<path>``.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import quote, unquote

NOTE_TYPES = {
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".rst": "text",
    ".org": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".csv": "csv",
    ".html": "html",
    ".htm": "html",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp"}
BROWSER_IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "webp"}
MIN_IMAGE_PX = 64  # skip icons, bullets, spacer images embedded in documents
MAX_TABLE_ROWS = 500
SCANNED_PAGE_CHARS = 25  # a PDF page with less extractable text than this is treated as a scan
OCR_DPI = 150


@dataclass
class ExtractedDoc:
    title: str
    markdown: str
    kind: str = "markdown"  # how the previewer should render `markdown`: markdown | text
    tags: list[str] = field(default_factory=list)


@dataclass
class Context:
    """Where a file lives and where to put derived assets."""

    root: Path
    cache_dir: Path
    ocr: Callable[[bytes], str] | None = None  # used for scanned PDF pages
    _names: dict[str, list[str]] | None = None

    def rel(self, p: Path) -> str:
        return p.resolve().relative_to(self.root.resolve()).as_posix()

    def find_by_name(self, name: str) -> str | None:
        """Obsidian-style lookup: resolve a bare file name anywhere in the vault."""
        if self._names is None:
            self._names = {}
            for p in self.root.rglob("*"):
                if p.is_file() and not any(part.startswith(".") for part in p.relative_to(self.root).parts):
                    rel = p.relative_to(self.root).as_posix()
                    self._names.setdefault(p.name.lower(), []).append(rel)
                    self._names.setdefault(p.stem.lower(), []).append(rel)
        hits = self._names.get(name.lower()) or self._names.get(Path(name).name.lower())
        return sorted(hits, key=len)[0] if hits else None


def file_url(rel: str) -> str:
    return "/api/file?path=" + quote(rel)


def doc_link(rel: str) -> str:
    return "#doc=" + quote(rel)


def _cache_image(ctx: Context, src: Path, key: str, data: bytes, ext: str) -> str:
    ext = ext.lower().lstrip(".")
    st = src.stat()
    name = hashlib.sha1(f"{src}:{st.st_mtime_ns}:{key}".encode()).hexdigest()[:20] + "." + ext
    out = ctx.cache_dir / name
    if not out.exists():
        ctx.cache_dir.mkdir(parents=True, exist_ok=True)
        out.write_bytes(data)
    return "/api/cache-image/" + name


def _image_size_ok(data: bytes) -> bool:
    try:
        import pymupdf as fitz

        pix = fitz.Pixmap(data)
        return pix.width >= MIN_IMAGE_PX and pix.height >= MIN_IMAGE_PX
    except Exception:
        return True  # unknown format (e.g. EMF) — keep and let the caller decide


def _to_browser_image(data: bytes, ext: str) -> tuple[bytes, str] | None:
    """Convert formats browsers can't show (tiff, jpx, emf…) to PNG when possible."""
    ext = ext.lower().lstrip(".")
    if ext == "jpeg":
        ext = "jpg"
    if ext in BROWSER_IMAGE_EXTS:
        return data, ext
    try:
        import pymupdf as fitz

        pix = fitz.Pixmap(data)
        if pix.n - pix.alpha >= 4:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        return pix.tobytes("png"), "png"
    except Exception:
        return None


def _md_escape_cell(s: str) -> str:
    return " ".join(str(s).split()).replace("|", "\\|")


def _md_table(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [[_md_escape_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# --------------------------------------------------------------------------- Markdown

FRONTMATTER_RE = re.compile(r"\A---\s*\n(.*?)\n(?:---|\.\.\.)\s*\n", re.S)
MD_IMAGE_RE = re.compile(r'!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+["\'][^"\']*["\'])?\s*\)')
WIKI_EMBED_RE = re.compile(r"!\[\[([^\]|#]+)(?:#[^\]|]*)?(?:\|([^\]]*))?\]\]")
WIKI_LINK_RE = re.compile(r"(?<!!)\[\[([^\]|#]+)(?:#[^\]|]*)?(?:\|([^\]]*))?\]\]")
MD_LINK_RE = re.compile(r"(?<!!)\[([^\]]*)\]\(\s*<?([^)\s>]+)>?\s*\)")
HTML_IMG_RE = re.compile(r'(<img\b[^>]*?\bsrc\s*=\s*)(["\'])(.*?)\2', re.I)
FENCE_RE = re.compile(r"^(```|~~~)", re.M)
EXTERNAL_RE = re.compile(r"^(?:[a-z][a-z0-9+.-]*:|//|#)", re.I)


def _resolve_local(ctx: Context, base: Path, target: str) -> str | None:
    """Resolve a link target written in a note to a path relative to the notes root."""
    target = unquote(target.split("#")[0].split("?")[0])
    if not target:
        return None
    root = ctx.root.resolve()
    cand = (root / target.lstrip("/")) if target.startswith("/") else (base.parent / target)
    try:
        cand = cand.resolve()
        cand.relative_to(root)
    except (ValueError, OSError):
        return None
    if cand.is_file():
        return cand.relative_to(root).as_posix()
    return ctx.find_by_name(Path(target).name)


def _split_code(md: str) -> list[tuple[bool, str]]:
    """Split Markdown into (is_code, text) segments so link rewriting skips fenced code."""
    out, pos, in_code, fence = [], 0, False, ""
    for m in FENCE_RE.finditer(md):
        line_end = md.find("\n", m.start())
        line_end = len(md) if line_end == -1 else line_end + 1
        if not in_code:
            out.append((False, md[pos : m.start()]))
            pos, in_code, fence = m.start(), True, m.group(1)
        elif m.group(1) == fence:
            out.append((True, md[pos:line_end]))
            pos, in_code = line_end, False
    out.append((in_code, md[pos:]))
    return out


def rewrite_markdown_links(md: str, ctx: Context, path: Path) -> str:
    def img(m: re.Match) -> str:
        alt, target = m.group(1), m.group(2)
        if EXTERNAL_RE.match(target):
            return m.group(0)
        rel = _resolve_local(ctx, path, target)
        return f"![{alt}]({file_url(rel)})" if rel else m.group(0)

    def wiki_embed(m: re.Match) -> str:
        target, alias = m.group(1).strip(), (m.group(2) or "").strip()
        rel = _resolve_local(ctx, path, target) or ctx.find_by_name(target)
        if not rel:
            return m.group(0)
        if Path(rel).suffix.lower() in IMAGE_EXTS:
            alt = "" if alias.isdigit() or re.fullmatch(r"\d+x\d+", alias) else alias
            return f"![{alt or Path(rel).stem}]({file_url(rel)})"
        return f"[{alias or Path(rel).stem}]({doc_link(rel)})"

    def wiki_link(m: re.Match) -> str:
        target, alias = m.group(1).strip(), (m.group(2) or "").strip()
        rel = ctx.find_by_name(target) or ctx.find_by_name(target + ".md")
        return f"[{alias or target}]({doc_link(rel)})" if rel else (alias or target)

    def link(m: re.Match) -> str:
        text, target = m.group(1), m.group(2)
        if EXTERNAL_RE.match(target):
            return m.group(0)
        rel = _resolve_local(ctx, path, target)
        if not rel:
            return m.group(0)
        suffix = Path(rel).suffix.lower()
        if suffix in NOTE_TYPES:
            return f"[{text}]({doc_link(rel)})"
        return f"[{text}]({file_url(rel)})"

    def html_img(m: re.Match) -> str:
        target = m.group(3)
        if EXTERNAL_RE.match(target):
            return m.group(0)
        rel = _resolve_local(ctx, path, target)
        return f"{m.group(1)}{m.group(2)}{file_url(rel)}{m.group(2)}" if rel else m.group(0)

    parts = []
    for is_code, seg in _split_code(md):
        if not is_code:
            seg = WIKI_EMBED_RE.sub(wiki_embed, seg)
            seg = MD_IMAGE_RE.sub(img, seg)
            seg = WIKI_LINK_RE.sub(wiki_link, seg)
            seg = MD_LINK_RE.sub(link, seg)
            seg = HTML_IMG_RE.sub(html_img, seg)
        parts.append(seg)
    return "".join(parts)


def _parse_frontmatter(md: str) -> tuple[dict, str]:
    m = FRONTMATTER_RE.match(md)
    if not m:
        return {}, md
    meta: dict = {}
    key = None
    for line in m.group(1).splitlines():
        if re.match(r"^\s+-\s+", line) and key:
            meta.setdefault(key, [])
            if isinstance(meta[key], list):
                meta[key].append(line.split("-", 1)[1].strip().strip("'\""))
            continue
        kv = re.match(r"^([A-Za-z_][\w-]*)\s*:\s*(.*)$", line)
        if kv:
            key, val = kv.group(1).lower(), kv.group(2).strip()
            if val.startswith("[") and val.endswith("]"):
                meta[key] = [v.strip().strip("'\"") for v in val[1:-1].split(",") if v.strip()]
            elif val:
                meta[key] = val.strip("'\"")
    return meta, md[m.end() :]


def extract_markdown(path: Path, ctx: Context) -> ExtractedDoc:
    raw = path.read_text(encoding="utf-8", errors="replace")
    meta, body = _parse_frontmatter(raw)
    tags = meta.get("tags", [])
    if isinstance(tags, str):
        tags = [t.strip() for t in re.split(r"[,\s]+", tags) if t.strip()]
    tags += [t for t in re.findall(r"(?:^|\s)#([A-Za-z][\w/-]+)", body) if t not in tags]
    title = meta.get("title")
    if not title:
        h1 = re.search(r"^#\s+(.+?)\s*#*\s*$", body, re.M)
        title = h1.group(1) if h1 else path.stem
    return ExtractedDoc(title=str(title), markdown=rewrite_markdown_links(body, ctx, path), tags=tags)


def extract_text(path: Path, ctx: Context) -> ExtractedDoc:
    return ExtractedDoc(title=path.stem, markdown=path.read_text(encoding="utf-8", errors="replace"), kind="text")


# --------------------------------------------------------------------------- PDF


def extract_pdf(path: Path, ctx: Context) -> ExtractedDoc:
    import pymupdf as fitz

    doc = fitz.open(path)
    try:
        title = (doc.metadata or {}).get("title") or path.stem
        parts: list[str] = []
        seen: set[int] = set()
        for pno, page in enumerate(doc, 1):
            parts.append(f"## Page {pno}\n")
            blocks = [" ".join(b[4].split()) for b in page.get_text("blocks", sort=True) if b[6] == 0]
            blocks = [b for b in blocks if b]
            scanned = sum(len(b) for b in blocks) < SCANNED_PAGE_CHARS and ctx.ocr is not None and page.get_images()
            if scanned:
                # No text layer: the page is (mostly) a picture, so read it with OCR instead.
                text = ctx.ocr(page.get_pixmap(dpi=OCR_DPI).tobytes("png"))
                if text:
                    parts.append("_Text recognized by OCR:_\n")
                    blocks = [line for line in text.splitlines() if line.strip()]
            for text in blocks:
                parts.append(re.sub(r"^([#>*+-]|\d+\.)", r"\\\1", text) + "\n")
            if scanned:
                continue  # the page scan itself isn't a useful figure; the PDF viewer shows it
            for info in page.get_images(full=True):
                xref, w, h = info[0], info[2], info[3]
                if xref in seen or w < MIN_IMAGE_PX or h < MIN_IMAGE_PX:
                    continue
                seen.add(xref)
                try:
                    img = doc.extract_image(xref)
                    conv = _to_browser_image(img["image"], img["ext"])
                    if not conv:
                        pix = fitz.Pixmap(doc, xref)
                        if pix.n - pix.alpha >= 4:
                            pix = fitz.Pixmap(fitz.csRGB, pix)
                        conv = (pix.tobytes("png"), "png")
                except Exception:
                    continue
                url = _cache_image(ctx, path, f"xref{xref}", *conv)
                parts.append(f"![Figure on page {pno}]({url})\n")
        return ExtractedDoc(title=title, markdown="\n".join(parts))
    finally:
        doc.close()


# --------------------------------------------------------------------------- DOCX


def extract_docx(path: Path, ctx: Context) -> ExtractedDoc:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    parts: list[str] = []
    img_count = 0

    def images_in(el) -> list[str]:
        nonlocal img_count
        out = []
        for blip in el.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = d.part.related_parts.get(rid) if rid else None
            if part is None or not hasattr(part, "blob"):
                continue
            data = part.blob
            conv = _to_browser_image(data, Path(str(part.partname)).suffix)
            if not conv or not _image_size_ok(conv[0]):
                continue
            img_count += 1
            out.append(f"![Image {img_count}]({_cache_image(ctx, path, rid + str(img_count), *conv)})")
        return out

    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            text = p.text.strip()
            style = (p.style.name if p.style is not None else "") or ""
            imgs = images_in(el)
            if text:
                m = re.match(r"Heading (\d)", style)
                if m:
                    parts.append("#" * min(int(m.group(1)), 6) + " " + text)
                elif style == "Title":
                    parts.append("# " + text)
                elif "List Number" in style:
                    parts.append("1. " + text)
                elif "List" in style or el.find(".//" + qn("w:numPr")) is not None:
                    parts.append("- " + text)
                else:
                    parts.append(text)
            parts.extend(imgs)
        elif el.tag == qn("w:tbl"):
            t = Table(el, d)
            rows = [[c.text for c in r.cells] for r in t.rows]
            parts.append(_md_table(rows))
            parts.extend(images_in(el))
    title = d.core_properties.title or path.stem
    return ExtractedDoc(title=title, markdown="\n\n".join(p for p in parts if p))


# --------------------------------------------------------------------------- PPTX


def extract_pptx(path: Path, ctx: Context) -> ExtractedDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    parts: list[str] = []
    img_count = 0

    def walk(shapes, out: list[str], title_shape_id):
        nonlocal img_count
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, out, title_shape_id)
                continue
            is_picture_placeholder = sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(sh, "image")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or is_picture_placeholder:
                try:
                    conv = _to_browser_image(sh.image.blob, sh.image.ext)
                except Exception:
                    conv = None
                if conv and _image_size_ok(conv[0]):
                    img_count += 1
                    alt = sh.name or f"Image {img_count}"
                    out.append(f"![{alt}]({_cache_image(ctx, path, f'img{img_count}', *conv)})")
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(_md_table([[c.text for c in r.cells] for r in sh.table.rows]))
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.shape_id != title_shape_id:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("  " * para.level + "- " + t)

    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        parts.append(f"## Slide {i}" + (f": {title}" if title else ""))
        body: list[str] = []
        walk(slide.shapes, body, title_shape.shape_id if title_shape is not None else None)
        parts.extend(body)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("**Speaker notes:** " + notes)
    title = prs.core_properties.title or path.stem
    return ExtractedDoc(title=title, markdown="\n\n".join(parts))


# --------------------------------------------------------------------------- spreadsheets


def extract_xlsx(path: Path, ctx: Context) -> ExtractedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if v is None else str(v) for v in row][:30])
                if len(rows) > MAX_TABLE_ROWS:
                    break
            parts.append(f"## Sheet: {ws.title}")
            parts.append(_md_table(rows) or "_(empty)_")
    finally:
        wb.close()
    return ExtractedDoc(title=path.stem, markdown="\n\n".join(parts))


def extract_csv(path: Path, ctx: Context) -> ExtractedDoc:
    text = path.read_text(encoding="utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append(row[:30])
        if len(rows) > MAX_TABLE_ROWS:
            break
    return ExtractedDoc(title=path.stem, markdown=_md_table(rows))


# --------------------------------------------------------------------------- HTML


class _HTMLToMarkdown(HTMLParser):
    BLOCK = {"p", "div", "section", "article", "br", "tr", "table", "ul", "ol", "blockquote", "pre"}
    SKIP = {"script", "style", "head", "noscript", "svg"}

    def __init__(self, on_img):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.on_img = on_img
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self.skip += 1
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.out.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in {"td", "th"}:
            self.out.append(" | ")
        elif tag in self.BLOCK:
            self.out.append("\n\n")
        elif tag == "img" and a.get("src"):
            md = self.on_img(a["src"], a.get("alt") or "")
            if md:
                self.out.append(f"\n\n{md}\n\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"} or tag in self.BLOCK:
            self.out.append("\n\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if not self.skip:
            self.out.append(re.sub(r"\s+", " ", data))


def extract_html(path: Path, ctx: Context) -> ExtractedDoc:
    def on_img(src: str, alt: str) -> str:
        if EXTERNAL_RE.match(src):
            return f"![{alt}]({src})" if src.startswith("http") else ""
        rel = _resolve_local(ctx, path, src)
        return f"![{alt}]({file_url(rel)})" if rel else ""

    p = _HTMLToMarkdown(on_img)
    p.feed(path.read_text(encoding="utf-8", errors="replace"))
    md = re.sub(r"\n{3,}", "\n\n", "".join(p.out))
    md = "\n".join(line.strip() for line in md.splitlines())
    return ExtractedDoc(title=p.title.strip() or path.stem, markdown=md.strip())


# ---------------------------------------------------------------------------

EXTRACTORS = {
    "markdown": extract_markdown,
    "text": extract_text,
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "xlsx": extract_xlsx,
    "csv": extract_csv,
    "html": extract_html,
}


def is_note(path: Path) -> bool:
    return path.suffix.lower() in NOTE_TYPES


def extract(path: Path, ctx: Context) -> ExtractedDoc:
    kind = NOTE_TYPES.get(path.suffix.lower())
    if not kind:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return EXTRACTORS[kind](path, ctx)

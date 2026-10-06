"""PowerPoint: one section per slide with its text, tables, pictures and speaker notes."""

from __future__ import annotations

from pathlib import Path

from .base import Context, ExtractedDoc, cache_image, image_size_ok, md_table, to_browser_image

def extract_pptx(path: Path, ctx: Context) -> ExtractedDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    parts: list[str] = []
    img_count = 0

    def walk(shapes, out: list[str], title_shape_id):
        nonlocal img_count
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, out, title_shape_id)
                continue
            is_picture_placeholder = sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(sh, "image")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or is_picture_placeholder:
                try:
                    conv = to_browser_image(sh.image.blob, sh.image.ext)
                except Exception:
                    conv = None
                if conv and image_size_ok(conv[0]):
                    img_count += 1
                    alt = sh.name or f"Image {img_count}"
                    out.append(f"![{alt}]({cache_image(ctx, path, f'img{img_count}', *conv)})")
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(md_table([[c.text for c in r.cells] for r in sh.table.rows]))
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.shape_id != title_shape_id:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("  " * para.level + "- " + t)

    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        parts.append(f"## Slide {i}" + (f": {title}" if title else ""))
        body: list[str] = []
        walk(slide.shapes, body, title_shape.shape_id if title_shape is not None else None)
        parts.extend(body)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("**Speaker notes:** " + notes)
    title = prs.core_properties.title or path.stem
    return ExtractedDoc(title=title, markdown="\n\n".join(parts))

"""Builds a sample notes directory containing every supported file type, with images."""

from __future__ import annotations

from pathlib import Path

import pytest


def make_png(w: int = 160, h: int = 120, rgb=(200, 40, 40)) -> bytes:
    import pymupdf as fitz

    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, w, h), 0)
    pix.set_rect(pix.irect, rgb)
    return pix.tobytes("png")


def build_notes(root: Path) -> Path:
    import docx
    import pymupdf as fitz
    import openpyxl
    from docx.shared import Inches
    from pptx import Presentation
    from pptx.util import Inches as PInches

    root.mkdir(parents=True, exist_ok=True)
    (root / "images").mkdir()
    (root / "images" / "vlan-diagram.png").write_bytes(make_png())
    (root / "networking").mkdir()
    (root / "networking" / "topology.png").write_bytes(make_png(rgb=(30, 90, 200)))

    (root / "networking" / "vlans.md").write_text(
        """---
title: VLAN Setup Guide
tags: [networking, homelab]
---
# VLAN Setup Guide

Intro paragraph about virtual LANs and trunk ports.

## Trunk configuration

Configure the switch trunk port with 802.1Q tagging for VLAN 10 and VLAN 20.

![VLAN diagram](../images/vlan-diagram.png)

## Topology

The topology is shown below.

![[topology.png]]

See also [[backups]] and [the router notes](router.md).

```bash
# not a heading
echo "![fake](nothere.png)"
```
"""
    )
    (root / "networking" / "router.md").write_text("# Router\n\nThe router runs OpenWrt with firewall zones.\n")
    (root / "backups.md").write_text(
        "# Backups\n\nRestic snapshots run nightly to the NAS. Retention keeps 7 daily and 4 weekly.\n"
    )
    (root / "todo.txt").write_text("Buy a new UPS battery.\nReplace the switch fan.\n")

    # PDF with text on two pages and an embedded image
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Kubernetes cluster upgrade procedure", fontsize=14)
    page.insert_text((72, 100), "Drain each node before upgrading kubelet.")
    page.insert_image(fitz.Rect(72, 150, 232, 270), stream=make_png(rgb=(20, 160, 60)))
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Rollback: restore the etcd snapshot taken before the upgrade.")
    pdf.set_metadata({"title": "K8s Upgrade"})
    pdf.save(root / "k8s-upgrade.pdf")
    pdf.close()

    # Word document with headings, list, table and image
    d = docx.Document()
    d.core_properties.title = "Onboarding Handbook"
    d.add_heading("Onboarding Handbook", 0)
    d.add_heading("Laptop setup", 1)
    d.add_paragraph("Install the VPN client and enroll in disk encryption.")
    d.add_paragraph("Request Git access", style="List Bullet")
    img_path = root / "_tmp.png"
    img_path.write_bytes(make_png(rgb=(120, 120, 20)))
    d.add_picture(str(img_path), width=Inches(2))
    img_path.unlink()
    d.add_heading("Contacts", 1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Team", "Email"
    t.cell(1, 0).text, t.cell(1, 1).text = "IT helpdesk", "it@example.com"
    d.save(root / "onboarding.docx")

    # PowerPoint with a title, bullets, picture and speaker notes
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Quarterly roadmap"
    tb = s.shapes.add_textbox(PInches(1), PInches(2), PInches(6), PInches(1))
    tb.text_frame.text = "Migrate monitoring to Prometheus"
    p2 = root / "_tmp2.png"
    p2.write_bytes(make_png(rgb=(90, 20, 140)))
    s.shapes.add_picture(str(p2), PInches(1), PInches(3))
    p2.unlink()
    s.notes_slide.notes_text_frame.text = "Mention Grafana dashboards."
    prs.save(root / "roadmap.pptx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(["Host", "IP", "Role"])
    ws.append(["nas01", "10.0.10.5", "storage"])
    ws.append(["pve01", "10.0.10.6", "hypervisor"])
    wb.save(root / "inventory.xlsx")

    (root / "subnets.csv").write_text("vlan,subnet,purpose\n10,10.0.10.0/24,servers\n20,10.0.20.0/24,iot\n")
    (root / "wiki.html").write_text(
        "<html><head><title>Wiki Export</title><script>alert(1)</script></head><body>"
        "<h1>DNS</h1><p>Pi-hole serves DNS for the lab.</p><img src='images/vlan-diagram.png' alt='dns pic'>"
        "<ul><li>Upstream: Quad9</li></ul></body></html>"
    )
    (root / "broken.pdf").write_bytes(b"this is not a pdf")
    (root / ".hidden").mkdir()
    (root / ".hidden" / "secret.md").write_text("# Secret\n\nshould not be indexed")
    return root


@pytest.fixture()
def notes(tmp_path) -> Path:
    return build_notes(tmp_path / "notes")


@pytest.fixture()
def data(tmp_path, monkeypatch) -> Path:
    d = tmp_path / "data"
    monkeypatch.setenv("RECALL_DATA_DIR", str(d))
    # share downloaded models across test runs instead of re-downloading per temp dir
    monkeypatch.setenv("RECALL_MODEL_DIR", str(Path.home() / ".cache" / "recall-test-models"))
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("RECALL_NOTES_DIR", raising=False)
    return d


@pytest.fixture()
def index(notes, data):
    from recall.index import Index

    idx = Index(notes, data)
    idx.build()
    return idx


def make_text_png(text: str, w: int = 520, h: int = 140) -> bytes:
    """A white image with black text on it (for OCR tests)."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=w, height=h)
    page.insert_text((20, h / 2 + 8), text, fontsize=22)
    return page.get_pixmap(dpi=150).tobytes("png")


def add_image_notes(root: Path) -> Path:
    """Extras for OCR/caption tests: a note with a text screenshot, a loose screenshot, a scanned PDF."""
    import pymupdf

    (root / "firewall").mkdir()
    (root / "firewall" / "rules.png").write_bytes(make_text_png("Allow TCP 8443 from DMZ"))
    (root / "firewall" / "edge.md").write_text("# Edge firewall\n\nCurrent rule set:\n\n![rules](rules.png)\n")
    (root / "whiteboard.png").write_bytes(make_text_png("Quarterly OKR brainstorm"))
    scan = pymupdf.open()
    page = scan.new_page(width=600, height=300)
    page.insert_image(page.rect, stream=make_text_png("Invoice total 4,250 EUR due March", 600, 300))
    scan.save(root / "scanned-invoice.pdf")
    scan.close()
    return root


@pytest.fixture()
def image_notes(notes) -> Path:
    return add_image_notes(notes)


@pytest.fixture(scope="session")
def embedder():
    """The real local embedding model; tests using it are skipped if it can't be loaded."""
    import os

    os.environ.setdefault("RECALL_MODEL_DIR", str(Path.home() / ".cache" / "recall-test-models"))
    from recall.embeddings import FastEmbedEmbedder, available

    if not available():
        pytest.skip("fastembed not installed")
    e = FastEmbedEmbedder()
    try:
        e.embed_query("warm up")
    except Exception as exc:
        pytest.skip(f"embedding model unavailable: {exc}")
    return e
